"""Extract text from PDF / DOCX / TXT resumes and heuristically structure it.

Parsing is best-effort: resumes have no fixed format, so the result is always
shown to the user for review and correction before it is saved.
"""

from __future__ import annotations

import io
import re
from typing import Dict, List, Optional, Tuple

from pydantic import BaseModel, Field

from models.resume_models import (
    Achievement,
    Certification,
    Education,
    Experience,
    MasterResume,
    PersonalInfo,
    Project,
    SkillSet,
)
from utils.file_utils import validate_upload
from utils.helpers import clean_text, split_items, strip_bullet, unique
from utils.logger import get_logger
from utils.skill_dictionary import extract_skills, normalize_skill, skill_category

logger = get_logger("parser")


class ParseError(Exception):
    """Raised with a user-friendly message when text extraction fails."""


class ParseResult(BaseModel):
    resume: MasterResume
    raw_text: str = ""
    sections_found: List[str] = Field(default_factory=list)
    warnings: List[str] = Field(default_factory=list)


SECTION_ALIASES: Dict[str, List[str]] = {
    "summary": [
        "summary", "professional summary", "profile", "professional profile", "about me",
        "objective", "career objective", "career summary", "executive summary", "overview",
        "profile summary",
    ],
    "skills": [
        "skills", "technical skills", "core skills", "key skills", "technologies",
        "technical expertise", "core competencies", "competencies", "skill set", "skillset",
        "skills and tools", "tools and technologies", "tech stack", "areas of expertise",
        "technical proficiency", "it skills", "skills summary",
    ],
    "experience": [
        "experience", "work experience", "professional experience", "employment",
        "employment history", "career history", "work history", "relevant experience",
        "internships", "internship", "internship experience", "professional background",
    ],
    "projects": [
        "projects", "personal projects", "academic projects", "key projects",
        "project experience", "selected projects", "notable projects", "project work",
    ],
    "education": [
        "education", "academic background", "educational qualifications",
        "academic qualifications", "qualifications", "education and training", "academics",
        "educational background",
    ],
    "certifications": [
        "certifications", "certificates", "licenses and certifications",
        "licences and certifications", "certification", "courses",
        "courses and certifications", "certifications and courses", "training",
    ],
    "achievements": [
        "achievements", "awards", "honors", "honours", "awards and achievements",
        "accomplishments", "awards and honors", "achievements and awards",
    ],
}

_HEADING_LOOKUP = {alias: key for key, aliases in SECTION_ALIASES.items() for alias in aliases}

_MONTH = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?"
_DATE = rf"(?:{_MONTH}\s*'?,?\s*\d{{2,4}}|\d{{1,2}}/\d{{4}}|\d{{4}})"
DATE_RANGE_RE = re.compile(
    rf"(?P<start>{_DATE})\s*(?:-|–|—|to|until|till)\s*"
    rf"(?P<end>{_DATE}|present|current|now|till date|to date|ongoing)",
    re.IGNORECASE,
)
SINGLE_DATE_RE = re.compile(rf"\b{_DATE}\b", re.IGNORECASE)

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
PHONE_RE = re.compile(r"(?<!\w)(\+?\d[\d\s().-]{7,}\d)(?!\w)")
LINKEDIN_RE = re.compile(r"(?:https?://)?(?:[a-z]{2,3}\.)?linkedin\.com/[^\s|,]+", re.IGNORECASE)
GITHUB_RE = re.compile(r"(?:https?://)?(?:www\.)?github\.com/[^\s|,]+", re.IGNORECASE)
URL_RE = re.compile(r"(?:https?://|www\.)[^\s|,]+", re.IGNORECASE)

TECH_LINE_RE = re.compile(
    r"^(?:technologies|technology|tech stack|tech|tools|tools used|environment|stack|skills used"
    r"|built with)\s*[:\-–]\s*(.+)$",
    re.IGNORECASE,
)
RESULT_LINE_RE = re.compile(r"^(?:results?|outcomes?|impact)\s*[:\-–]\s*(.+)$", re.IGNORECASE)

TITLE_WORDS_RE = re.compile(
    r"\b(engineer|developer|manager|analyst|intern|scientist|lead|consultant|architect|designer"
    r"|specialist|associate|administrator|officer|director|head|executive|programmer|trainee"
    r"|researcher|assistant|coordinator|tester|technician|founder|freelancer?|apprentice)\b",
    re.IGNORECASE,
)
COMPANY_WORDS_RE = re.compile(
    r"\b(ltd|limited|inc|corp|corporation|llc|llp|pvt|private|technologies|technology|solutions"
    r"|systems|labs|group|software|services|consulting|company|university|bank|gmbh|plc"
    r"|enterprises|industries|infotech|global)\b\.?",
    re.IGNORECASE,
)
LEGAL_SUFFIX_RE = re.compile(
    r"\b(ltd|limited|inc|corp|corporation|llc|llp|pvt|private|gmbh|plc|co\.)\b\.?", re.IGNORECASE
)
DEGREE_RE = re.compile(
    r"\b(b\.?\s?tech|m\.?\s?tech|b\.\s?e\b|m\.\s?e\b|b\.?\s?sc|m\.?\s?sc|bca|mca|mba|bba|b\.?\s?com"
    r"|m\.?\s?com|ph\.?\s?d|bachelor'?s?|master'?s?|diploma|degree|hsc|ssc|12th|10th"
    r"|higher secondary|secondary school|intermediate|pre-university|puc)\b",
    re.IGNORECASE,
)
DEGREE_CS_RE = re.compile(r"\b(BE|ME|BS|MS|BA|MA|BSc|MSc)\b")
INSTITUTION_RE = re.compile(
    r"\b(university|college|institute|school|academy|iit|nit|iiit|polytechnic|vidyalaya"
    r"|vidyapeeth)\b",
    re.IGNORECASE,
)
GPA_RE = re.compile(
    r"(?:cgpa|gpa|percentage|score|grade|marks)\s*[:\-]?\s*(\d+(?:\.\d+)?\s*(?:/\s*\d+(?:\.\d+)?)?\s*%?)"
    r"|(\d{2}(?:\.\d+)?\s*%)",
    re.IGNORECASE,
)

CATEGORY_LABELS: Dict[str, List[str]] = {
    "deep_learning": ["deep learning", "dl"],
    "nlp": ["nlp", "natural language"],
    "ai_ml": ["machine learning", "ai", "ml", "data science", "artificial intelligence", "ai/ml"],
    "programming_languages": ["language", "programming"],
    "frameworks": ["framework", "librar"],
    "databases": ["database", "db"],
    "cloud": ["cloud"],
    "devops": ["devops", "ci/cd"],
    "web_development": ["web", "front", "back", "ui"],
    "tools": ["tool", "software", "platform", "ide"],
    "other": ["other", "soft", "misc", "methodolog", "concepts"],
}


# --------------------------------------------------------------------------- text extraction

def extract_text_from_pdf(data: bytes) -> str:
    import pymupdf

    try:
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            if doc.needs_pass:
                raise ParseError("This PDF is password protected. Please upload an unlocked copy.")
            return "\n".join(page.get_text("text") for page in doc)
    except ParseError:
        raise
    except Exception as exc:
        logger.warning("PDF extraction failed: %s", type(exc).__name__)
        raise ParseError("Could not read this PDF. It may be corrupted.") from exc


def extract_text_from_docx(data: bytes) -> str:
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        logger.warning("DOCX extraction failed: %s", type(exc).__name__)
        raise ParseError("Could not read this DOCX file. It may be corrupted.") from exc
    lines: List[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text
        style = (paragraph.style.name if paragraph.style is not None else "").lower()
        if text.strip() and "list" in style and not strip_bullet(text)[0]:
            text = "• " + text
        lines.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = unique(cell.text.strip() for cell in row.cells)
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def extract_text_from_txt(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            if "\x00" not in text:
                return text
        except UnicodeDecodeError:
            continue
    raise ParseError("Could not decode this text file.")


def extract_text(data: bytes, filename: str) -> str:
    ext = validate_upload(filename, data)
    if ext == ".pdf":
        text = extract_text_from_pdf(data)
    elif ext == ".docx":
        text = extract_text_from_docx(data)
    else:
        text = extract_text_from_txt(data)
    text = clean_text(text)
    if len(text) < 30:
        raise ParseError(
            "Very little text could be extracted. If this is a scanned PDF, please upload a "
            "text-based PDF, DOCX or TXT version."
        )
    return text


# --------------------------------------------------------------------------- section detection

def _normalize_heading(line: str) -> str:
    value = line.strip().strip("#*_=:-–—|•").strip().lower()
    value = value.replace("&", " and ").replace("/", " and ")
    return re.sub(r"\s+", " ", value)


def detect_heading(line: str) -> Optional[str]:
    if not line or len(line) > 60 or len(line.split()) > 6:
        return None
    return _HEADING_LOOKUP.get(_normalize_heading(line))


def detect_sections(text: str) -> Dict[str, List[str]]:
    """Split resume text into ``{section_key: lines}``; preamble goes to ``header``."""
    sections: Dict[str, List[str]] = {"header": []}
    current = "header"
    for line in text.split("\n"):
        heading = detect_heading(line)
        if heading:
            current = heading
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)
    return sections


# --------------------------------------------------------------------------- helpers

def _has_date(text: str) -> bool:
    return bool(DATE_RANGE_RE.search(text))


def _is_long(text: str) -> bool:
    return len(text.split()) > 12


def _looks_like_header(text: str) -> bool:
    return bool(
        _has_date(text) or "|" in text or TITLE_WORDS_RE.search(text)
        or COMPANY_WORDS_RE.search(text) or DEGREE_RE.search(text)
    ) and len(text.split()) <= 14


def _is_continuation(previous: str, text: str) -> bool:
    if not previous or not text:
        return False
    if text[0].islower() or text[0] in "(&":
        return True
    if previous.rstrip().endswith((",", "-", "&", " and", " or", " with", " of", " the", " for")):
        return True
    return len(previous) >= 75 and not previous.rstrip().endswith((".", "!", "?", ":")) \
        and not _looks_like_header(text) and not TECH_LINE_RE.match(text)


def split_entries(lines: List[str]) -> List[Dict[str, List[str]]]:
    """Group section lines into entries of ``{"header": [...], "bullets": [...]}``."""
    entries: List[Dict[str, List[str]]] = []
    current: Optional[Dict[str, List[str]]] = None
    for raw in lines:
        if not raw.strip():
            continue
        is_bullet, text = strip_bullet(raw)
        if not text:
            continue
        if is_bullet:
            if current is None:
                current = {"header": [], "bullets": []}
                entries.append(current)
            current["bullets"].append(text)
            continue
        if current is not None:
            last_list = current["bullets"] or current["header"]
            if last_list and _is_continuation(last_list[-1], text):
                last_list[-1] = f"{last_list[-1]} {text}"
                continue
            if TECH_LINE_RE.match(text) or RESULT_LINE_RE.match(text):
                current["header"].append(text)
                continue
        header_has_date = current is not None and any(_has_date(h) for h in current["header"])
        starts_new = (
            current is None
            or bool(current["bullets"])
            or (current["header"] and _is_long(current["header"][-1]) and not _is_long(text))
            or (header_has_date and _has_date(text))
        )
        if starts_new:
            current = {"header": [], "bullets": []}
            entries.append(current)
        current["header"].append(text)
    return entries


def _split_segments(line: str) -> List[str]:
    parts = re.split(r"\s*[|•·]\s*|\s+[-–—]\s+|\s+at\s+|\s+@\s+", line)
    return [p.strip(" ,;()") for p in parts if p and p.strip(" ,;()")]


def _remove_date_range(line: str) -> Tuple[str, str, str]:
    match = DATE_RANGE_RE.search(line)
    if not match:
        return line, "", ""
    cleaned = (line[:match.start()] + " " + line[match.end():]).strip()
    cleaned = re.sub(r"\(\s*\)", "", cleaned)
    return cleaned, match.group("start").strip(), match.group("end").strip().title()


def _is_location(segment: str) -> bool:
    if segment.lower() in {"remote", "hybrid", "onsite", "on-site"}:
        return True
    return bool(re.fullmatch(r"[A-Z][A-Za-z .'-]+,\s*[A-Z][A-Za-z .'-]+", segment)) \
        and not COMPANY_WORDS_RE.search(segment) and not TITLE_WORDS_RE.search(segment)


# --------------------------------------------------------------------------- section parsers

def parse_personal(header_lines: List[str], full_text: str) -> PersonalInfo:
    info = PersonalInfo()
    head_text = "\n".join(header_lines) if header_lines else "\n".join(full_text.split("\n")[:8])
    search_text = head_text + "\n" + full_text[:1500]

    if m := EMAIL_RE.search(search_text):
        info.email = m.group(0)
    for m in PHONE_RE.finditer(search_text):
        digits = re.sub(r"\D", "", m.group(1))
        if 10 <= len(digits) <= 15 and not DATE_RANGE_RE.search(m.group(1)):
            info.phone = m.group(1).strip()
            break
    if m := LINKEDIN_RE.search(search_text):
        info.linkedin = m.group(0).rstrip(".")
    if m := GITHUB_RE.search(search_text):
        info.github = m.group(0).rstrip(".")
    for m in URL_RE.finditer(search_text):
        url = m.group(0).rstrip(".")
        if "linkedin.com" not in url.lower() and "github.com" not in url.lower():
            info.portfolio = url
            break

    candidates = header_lines or full_text.split("\n")[:5]
    for line in candidates:
        text = line.strip()
        if not text or EMAIL_RE.search(text) or URL_RE.search(text) or re.search(r"\d{3}", text):
            continue
        words = text.split()
        if 1 < len(words) <= 5 and all(re.fullmatch(r"[A-Za-z.'-]+", w) for w in words):
            info.name = text.title() if text.isupper() else text
            break

    for line in header_lines[:8]:
        for segment in re.split(r"\s*[|•·]\s*", line):
            segment = segment.strip()
            if segment and segment != info.name and _is_location(segment):
                info.location = segment
                break
        if info.location:
            break
    return info


def _category_for_label(label: str) -> Optional[str]:
    label = label.strip().lower()
    for category, keys in CATEGORY_LABELS.items():
        for key in keys:
            if len(key) <= 3:
                if re.search(rf"(?<![a-z]){re.escape(key)}(?![a-z])", label):
                    return category
            elif key in label:
                return category
    return None


def parse_skills(lines: List[str]) -> SkillSet:
    skills = SkillSet()
    for raw in lines:
        _, line = strip_bullet(raw)
        if not line:
            continue
        category: Optional[str] = None
        content = line
        label_match = re.match(r"^([^:]{2,40}):\s*(.+)$", line)
        if label_match:
            category = _category_for_label(label_match.group(1))
            content = label_match.group(2)
        items = split_items(content)
        if len(items) <= 1 and _is_long(content):
            items = extract_skills(content)
        for item in items:
            if not item or len(item) > 50:
                continue
            canonical = normalize_skill(item)
            target = category or (skill_category(canonical) if canonical else "other")
            skills.add(target, canonical or item)
    return skills


def parse_experience(lines: List[str]) -> List[Experience]:
    experiences: List[Experience] = []
    for entry in split_entries(lines):
        exp = Experience()
        header_segments: List[str] = []
        description: List[str] = []
        technologies: List[str] = []
        for line in entry["header"]:
            if m := TECH_LINE_RE.match(line):
                technologies.extend(split_items(m.group(1)))
                continue
            line, start, end = _remove_date_range(line)
            if start and not exp.start_date:
                exp.start_date, exp.end_date = start, end
            line = line.strip(" |,-–—")
            if not line:
                continue
            if _is_long(line):
                description.append(line)
            else:
                header_segments.extend(_split_segments(line))
        bullets: List[str] = []
        for bullet in entry["bullets"]:
            if m := TECH_LINE_RE.match(bullet):
                technologies.extend(split_items(m.group(1)))
            else:
                bullets.append(bullet)

        unassigned: List[str] = []
        for segment in header_segments:
            is_title = bool(TITLE_WORDS_RE.search(segment)) and not LEGAL_SUFFIX_RE.search(segment)
            if not exp.location and _is_location(segment):
                exp.location = segment
            elif not exp.title and is_title:
                exp.title = segment
            elif not exp.company and COMPANY_WORDS_RE.search(segment) and not is_title:
                exp.company = segment
            else:
                unassigned.append(segment)
        for segment in unassigned:
            if not exp.title:
                exp.title = segment
            elif not exp.company:
                exp.company = segment
            elif not exp.location and len(segment.split()) <= 3:
                exp.location = segment
            else:
                description.append(segment)

        exp.description = " ".join(description)
        exp.bullets = bullets
        exp.technologies = unique([normalize_skill(t) or t for t in technologies]
                                  + extract_skills(exp.text()))
        if exp.title or exp.company or exp.bullets:
            experiences.append(exp)
    return experiences


def parse_projects(lines: List[str]) -> List[Project]:
    projects: List[Project] = []
    for entry in split_entries(lines):
        project = Project()
        technologies: List[str] = []
        description: List[str] = []
        for index, line in enumerate(entry["header"]):
            if m := TECH_LINE_RE.match(line):
                technologies.extend(split_items(m.group(1)))
                continue
            if m := RESULT_LINE_RE.match(line):
                project.results = m.group(1).strip()
                continue
            line, _, _ = _remove_date_range(line)
            line = line.strip(" |,-–—")
            if not line:
                continue
            if index == 0 and not project.name and not _is_long(line):
                paren = re.search(r"\(([^)]*)\)", line)
                if paren and extract_skills(paren.group(1)):
                    technologies.extend(split_items(paren.group(1)))
                    line = (line[:paren.start()] + line[paren.end():]).strip()
                segments = _split_segments(line)
                project.name = segments[0] if segments else line
                for extra in segments[1:]:
                    if extract_skills(extra):
                        technologies.extend(split_items(extra))
                    else:
                        description.append(extra)
            else:
                description.append(line)
        responsibilities: List[str] = []
        for bullet in entry["bullets"]:
            if m := TECH_LINE_RE.match(bullet):
                technologies.extend(split_items(m.group(1)))
            elif m := RESULT_LINE_RE.match(bullet):
                project.results = m.group(1).strip()
            else:
                responsibilities.append(bullet)
        if not project.name and responsibilities:
            project.name = responsibilities.pop(0)[:80]
        project.description = " ".join(description)
        project.responsibilities = responsibilities
        project.technologies = unique([normalize_skill(t) or t for t in technologies]
                                      + extract_skills(project.text()))
        if project.name:
            projects.append(project)
    return projects


def parse_education(lines: List[str]) -> List[Education]:
    groups: List[List[str]] = []
    for raw in lines:
        _, line = strip_bullet(raw)
        if not line:
            continue
        has_degree = bool(DEGREE_RE.search(line) or DEGREE_CS_RE.search(line))
        current_has_degree = bool(groups) and any(
            DEGREE_RE.search(item) or DEGREE_CS_RE.search(item) for item in groups[-1]
        )
        if not groups or (has_degree and current_has_degree):
            groups.append([])
        groups[-1].append(line)

    education: List[Education] = []
    for group in groups:
        edu = Education()
        for line in group:
            if (m := GPA_RE.search(line)) and not edu.gpa:
                edu.gpa = (m.group(1) or m.group(2) or "").strip()
                line = (line[:m.start()] + line[m.end():]).strip()
            line, start, end = _remove_date_range(line)
            if start and not edu.start_date:
                edu.start_date, edu.end_date = start, end
            elif not edu.end_date and (m := SINGLE_DATE_RE.search(line)):
                edu.end_date = m.group(0)
                line = (line[:m.start()] + line[m.end():]).strip()
            for segment in re.split(r"\s*[|,•·]\s*|\s+[-–—]\s+", line):
                segment = segment.strip(" ()-–")
                if not segment:
                    continue
                if not edu.degree and (DEGREE_RE.search(segment) or DEGREE_CS_RE.search(segment)):
                    edu.degree = segment
                elif not edu.institution and INSTITUTION_RE.search(segment):
                    edu.institution = segment
                elif not edu.location and len(segment.split()) <= 3 and segment[:1].isupper() \
                        and (edu.degree or edu.institution):
                    edu.location = segment
                elif not edu.degree:
                    edu.degree = segment
                elif not edu.institution:
                    edu.institution = segment
        if edu.degree or edu.institution:
            education.append(edu)
    return education


def parse_certifications(lines: List[str]) -> List[Certification]:
    certifications: List[Certification] = []
    for raw in lines:
        _, line = strip_bullet(raw)
        if not line:
            continue
        credential_match = re.match(r"^(?:credential(?: id)?|id|url)\s*[:\-]\s*(.+)$", line, re.I)
        if credential_match and certifications:
            certifications[-1].credential = credential_match.group(1).strip()
            continue
        cert = Certification()
        if m := URL_RE.search(line):
            cert.credential = m.group(0)
            line = line.replace(m.group(0), "").strip()
        if m := SINGLE_DATE_RE.search(line):
            cert.date = m.group(0)
            line = (line[:m.start()] + line[m.end():]).strip()
        parts = [p.strip(" ,()-–") for p in re.split(r"\s*[|]\s*|\s+[-–—]\s+|\s+by\s+|,\s+", line)]
        parts = [p for p in parts if p]
        if not parts:
            continue
        cert.name = parts[0]
        if len(parts) > 1:
            cert.issuer = parts[1]
        certifications.append(cert)
    return certifications


def parse_achievements(lines: List[str]) -> List[Achievement]:
    achievements: List[Achievement] = []
    for raw in lines:
        _, line = strip_bullet(raw)
        if not line:
            continue
        if achievements and _is_continuation(achievements[-1].description or achievements[-1].title, line):
            target = achievements[-1]
            if target.description:
                target.description += " " + line
            else:
                target.title += " " + line
            continue
        title, _, description = line.partition(":")
        if description and len(title.split()) <= 8:
            achievements.append(Achievement(title=title.strip(), description=description.strip()))
        else:
            achievements.append(Achievement(title=line))
    return achievements


# --------------------------------------------------------------------------- public API

def parse_resume_text(text: str) -> ParseResult:
    text = clean_text(text)
    if not text:
        raise ParseError("The resume is empty.")
    sections = detect_sections(text)
    found = [key for key in sections if key != "header" and any(l.strip() for l in sections[key])]
    resume = MasterResume(
        personal=parse_personal([l for l in sections.get("header", []) if l.strip()], text),
        summary=" ".join(l.strip() for l in sections.get("summary", []) if l.strip()),
        skills=parse_skills(sections.get("skills", [])),
        experience=parse_experience(sections.get("experience", [])),
        projects=parse_projects(sections.get("projects", [])),
        education=parse_education(sections.get("education", [])),
        certifications=parse_certifications(sections.get("certifications", [])),
        achievements=parse_achievements(sections.get("achievements", [])),
    )
    warnings: List[str] = []
    if not found:
        warnings.append("No standard section headings were detected. Please fill in the sections manually.")
    if not resume.skills.all_skills():
        for skill in extract_skills(text):
            resume.skills.add(skill_category(skill), skill)
        if resume.skills.all_skills():
            warnings.append("No skills section found — skills were detected from the resume text. "
                            "Please review them.")
    if not resume.personal.name:
        warnings.append("Could not detect your name. Please enter it manually.")
    return ParseResult(resume=resume, raw_text=text, sections_found=found, warnings=warnings)


def parse_resume_file(data: bytes, filename: str) -> ParseResult:
    """Extract and structure a resume from uploaded bytes."""
    text = extract_text(data, filename)
    try:
        return parse_resume_text(text)
    except ParseError:
        raise
    except Exception as exc:
        logger.exception("Resume structuring failed")
        raise ParseError("The text was extracted but could not be structured. "
                         "Please fill in the Master Resume manually.") from exc

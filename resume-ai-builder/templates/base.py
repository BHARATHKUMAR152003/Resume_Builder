"""Shared resume template interface and rendering (no AI logic here)."""

from __future__ import annotations

import io
from abc import ABC
from dataclasses import dataclass
from typing import List, Tuple
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

from models.resume_models import TailoredResume

PDF_FONTS = {
    "Helvetica": ("Helvetica", "Helvetica-Bold", "Helvetica-Oblique"),
    "Times": ("Times-Roman", "Times-Bold", "Times-Italic"),
}


@dataclass(frozen=True)
class TemplateStyle:
    key: str
    label: str
    pdf_font: str = "Helvetica"
    docx_font: str = "Calibri"
    accent: str = "#000000"
    text: str = "#111111"
    muted: str = "#444444"
    name_size: float = 20
    heading_size: float = 11.5
    body_size: float = 10
    center_header: bool = True
    heading_upper: bool = True
    heading_rule: bool = True
    margin_mm: float = 18
    section_space: float = 8


def date_range(start: str, end: str) -> str:
    start, end = (start or "").strip(), (end or "").strip()
    if start and end:
        return f"{start} – {end}"
    return start or end


def resume_sections(resume: TailoredResume) -> List[Tuple[str, str]]:
    """Ordered list of (key, title) for sections that contain information."""
    sections = []
    if resume.summary.strip():
        sections.append(("summary", "Professional Summary"))
    if any(g.skills for g in resume.skills):
        sections.append(("skills", "Technical Skills"))
    if resume.experience:
        sections.append(("experience", "Professional Experience"))
    if resume.projects:
        sections.append(("projects", "Projects"))
    if resume.education:
        sections.append(("education", "Education"))
    if resume.certifications:
        sections.append(("certifications", "Certifications"))
    if resume.achievements:
        sections.append(("achievements", "Achievements"))
    return sections


class ResumeTemplate(ABC):
    """Common interface. Subclasses only provide a ``TemplateStyle``."""

    style: TemplateStyle

    @property
    def key(self) -> str:
        return self.style.key

    @property
    def label(self) -> str:
        return self.style.label

    def render(self, resume_data: TailoredResume, fmt: str = "pdf") -> bytes:
        if fmt == "pdf":
            return self.render_pdf(resume_data)
        if fmt == "docx":
            return self.render_docx(resume_data)
        raise ValueError(f"Unsupported format: {fmt}")

    # ------------------------------------------------------------------ PDF

    def _pdf_styles(self) -> dict:
        s = self.style
        regular, bold, italic = PDF_FONTS[s.pdf_font]
        align = TA_CENTER if s.center_header else TA_LEFT
        body = ParagraphStyle("body", fontName=regular, fontSize=s.body_size,
                              leading=s.body_size * 1.35, textColor=colors.HexColor(s.text))
        return {
            "name": ParagraphStyle("name", parent=body, fontName=bold, fontSize=s.name_size,
                                   leading=s.name_size * 1.2, alignment=align,
                                   textColor=colors.HexColor(s.accent), spaceAfter=2),
            "contact": ParagraphStyle("contact", parent=body, fontSize=s.body_size - 0.5,
                                      alignment=align, textColor=colors.HexColor(s.muted),
                                      spaceAfter=4),
            "heading": ParagraphStyle("heading", parent=body, fontName=bold,
                                      fontSize=s.heading_size, leading=s.heading_size * 1.3,
                                      textColor=colors.HexColor(s.accent),
                                      spaceBefore=s.section_space, spaceAfter=2,
                                      keepWithNext=1),
            "body": body,
            "item_title": ParagraphStyle("item_title", parent=body, fontName=bold, spaceBefore=4,
                                         keepWithNext=1),
            "meta": ParagraphStyle("meta", parent=body, fontName=italic, keepWithNext=1,
                                   fontSize=s.body_size - 0.5, textColor=colors.HexColor(s.muted)),
            "tech": ParagraphStyle("tech", parent=body, fontSize=s.body_size - 0.5,
                                   textColor=colors.HexColor(s.muted)),
            "bullet": ParagraphStyle("bullet", parent=body, leftIndent=12, bulletIndent=3,
                                     spaceBefore=1),
        }

    def _pdf_heading(self, title: str, st: dict) -> list:
        text = title.upper() if self.style.heading_upper else title
        flow: list = [Paragraph(escape(text), st["heading"])]
        if self.style.heading_rule:
            rule = HRFlowable(width="100%", thickness=0.7, spaceBefore=0, spaceAfter=3,
                              color=colors.HexColor(self.style.accent))
            rule.keepWithNext = True  # never leave a heading orphaned at a page bottom
            flow.append(rule)
        return flow

    @staticmethod
    def _bullets(items: List[str], st: dict) -> list:
        return [Paragraph(escape(i), st["bullet"], bulletText="•") for i in items if i.strip()]

    def _pdf_section(self, key: str, resume: TailoredResume, st: dict) -> list:
        flow: list = []
        if key == "summary":
            flow.append(Paragraph(escape(resume.summary), st["body"]))
        elif key == "skills":
            for group in resume.skills:
                if group.skills:
                    flow.append(Paragraph(f"<b>{escape(group.category)}:</b> "
                                          f"{escape(', '.join(group.skills))}", st["body"]))
        elif key == "experience":
            for exp in resume.experience:
                title = escape(exp.title or "")
                company = escape(exp.company or "")
                head = f"{title} — {company}" if title and company else (title or company)
                meta = " | ".join(p for p in (escape(exp.location),
                                              escape(date_range(exp.start_date, exp.end_date))) if p)
                block = [Paragraph(head, st["item_title"])]
                if meta:
                    block.append(Paragraph(meta, st["meta"]))
                if exp.description.strip():
                    block.append(Paragraph(escape(exp.description), st["body"]))
                flow.extend(block + self._bullets(exp.bullets, st))
                if exp.technologies:
                    flow.append(Paragraph(f"<i>Technologies:</i> {escape(', '.join(exp.technologies))}",
                                          st["tech"]))
        elif key == "projects":
            for project in resume.projects:
                block = [Paragraph(escape(project.name), st["item_title"])]
                if project.technologies:
                    block.append(Paragraph(escape(", ".join(project.technologies)), st["meta"]))
                if project.description.strip():
                    block.append(Paragraph(escape(project.description), st["body"]))
                bullets = self._bullets(project.responsibilities + project.achievements, st)
                if project.results.strip():
                    bullets += self._bullets([f"Result: {project.results}"], st)
                flow.extend(block + bullets)
        elif key == "education":
            for edu in resume.education:
                head = " — ".join(escape(p) for p in (edu.degree, edu.institution) if p)
                meta = " | ".join(escape(p) for p in (
                    edu.location, date_range(edu.start_date, edu.end_date),
                    f"GPA/Score: {edu.gpa}" if edu.gpa else "") if p)
                block = [Paragraph(head, st["item_title"])]
                if meta:
                    block.append(Paragraph(meta, st["meta"]))
                flow.extend(block)
        elif key == "certifications":
            items = []
            for cert in resume.certifications:
                text = " — ".join(x for x in (cert.name, cert.issuer) if x)
                if cert.date:
                    text += f" ({cert.date})"
                if cert.credential:
                    text += f" | {cert.credential}"
                items.append(text)
            flow.extend(self._bullets(items, st))
        elif key == "achievements":
            items = [f"{a.title}: {a.description}" if a.description else a.title
                     for a in resume.achievements]
            flow.extend(self._bullets(items, st))
        return flow

    def render_pdf(self, resume: TailoredResume) -> bytes:
        buffer = io.BytesIO()
        margin = self.style.margin_mm * mm
        name = resume.personal.name or "Resume"
        doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=margin, rightMargin=margin,
                                topMargin=margin * 0.8, bottomMargin=margin * 0.8,
                                title=f"{name} - Resume", author=name)
        st = self._pdf_styles()
        story: list = [Paragraph(escape(name), st["name"])]
        contacts = resume.personal.contact_items()
        if contacts:
            story.append(Paragraph("  |  ".join(escape(c) for c in contacts), st["contact"]))
        for key, title in resume_sections(resume):
            story.extend(self._pdf_heading(title, st))
            story.extend(self._pdf_section(key, resume, st))
        story.append(Spacer(1, 2))
        doc.build(story)
        return buffer.getvalue()

    # ------------------------------------------------------------------ DOCX

    @staticmethod
    def _rgb(hex_color: str) -> RGBColor:
        return RGBColor.from_string(hex_color.lstrip("#").upper())

    def _docx_run(self, paragraph, text: str, bold: bool = False, italic: bool = False,
                  size: float | None = None, color: str | None = None):
        run = paragraph.add_run(text)
        run.bold, run.italic = bold, italic
        run.font.name = self.style.docx_font
        run.font.size = Pt(size or self.style.body_size)
        run.font.color.rgb = self._rgb(color or self.style.text)
        return run

    def _docx_paragraph(self, doc, space_before: float = 0, space_after: float = 1,
                        center: bool = False, style: str | None = None):
        paragraph = doc.add_paragraph(style=style) if style else doc.add_paragraph()
        fmt = paragraph.paragraph_format
        fmt.space_before, fmt.space_after = Pt(space_before), Pt(space_after)
        if center:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return paragraph

    def _docx_heading(self, doc, title: str) -> None:
        paragraph = doc.add_paragraph()
        if self.style.heading_rule:
            p_pr = paragraph._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for attr, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"),
                                ("w:color", self.style.accent.lstrip("#"))):
                bottom.set(qn(attr), value)
            borders.append(bottom)
            p_pr.append(borders)
        fmt = paragraph.paragraph_format
        fmt.space_before, fmt.space_after = Pt(self.style.section_space + 2), Pt(3)
        fmt.keep_with_next = True
        text = title.upper() if self.style.heading_upper else title
        self._docx_run(paragraph, text, bold=True, size=self.style.heading_size,
                       color=self.style.accent)

    def _docx_bullet(self, doc, text: str) -> None:
        paragraph = self._docx_paragraph(doc, style="List Bullet", space_after=1)
        self._docx_run(paragraph, text)

    def render_docx(self, resume: TailoredResume) -> bytes:
        s = self.style
        doc = Document()
        section = doc.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.left_margin = section.right_margin = Mm(s.margin_mm)
        section.top_margin = section.bottom_margin = Mm(s.margin_mm * 0.8)
        normal = doc.styles["Normal"]
        normal.font.name = s.docx_font
        normal.font.size = Pt(s.body_size)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), s.docx_font)
        doc.core_properties.title = f"{resume.personal.name or 'Resume'} - Resume"
        doc.core_properties.author = resume.personal.name or ""

        p = self._docx_paragraph(doc, center=s.center_header, space_after=2)
        self._docx_run(p, resume.personal.name or "Resume", bold=True, size=s.name_size,
                       color=s.accent)
        contacts = resume.personal.contact_items()
        if contacts:
            p = self._docx_paragraph(doc, center=s.center_header, space_after=4)
            self._docx_run(p, "  |  ".join(contacts), size=s.body_size - 0.5, color=s.muted)

        for key, title in resume_sections(resume):
            self._docx_heading(doc, title)
            if key == "summary":
                self._docx_run(self._docx_paragraph(doc), resume.summary)
            elif key == "skills":
                for group in resume.skills:
                    if group.skills:
                        p = self._docx_paragraph(doc)
                        self._docx_run(p, f"{group.category}: ", bold=True)
                        self._docx_run(p, ", ".join(group.skills))
            elif key == "experience":
                for exp in resume.experience:
                    p = self._docx_paragraph(doc, space_before=4)
                    p.paragraph_format.keep_with_next = True
                    head = " — ".join(x for x in (exp.title, exp.company) if x)
                    self._docx_run(p, head, bold=True)
                    meta = " | ".join(x for x in (exp.location,
                                                  date_range(exp.start_date, exp.end_date)) if x)
                    if meta:
                        p = self._docx_paragraph(doc)
                        p.paragraph_format.keep_with_next = True
                        self._docx_run(p, meta, italic=True, size=s.body_size - 0.5, color=s.muted)
                    if exp.description.strip():
                        self._docx_run(self._docx_paragraph(doc), exp.description)
                    for bullet in exp.bullets:
                        self._docx_bullet(doc, bullet)
                    if exp.technologies:
                        p = self._docx_paragraph(doc)
                        self._docx_run(p, "Technologies: ", italic=True, size=s.body_size - 0.5,
                                       color=s.muted)
                        self._docx_run(p, ", ".join(exp.technologies), size=s.body_size - 0.5,
                                       color=s.muted)
            elif key == "projects":
                for project in resume.projects:
                    p = self._docx_paragraph(doc, space_before=4)
                    p.paragraph_format.keep_with_next = True
                    self._docx_run(p, project.name, bold=True)
                    if project.technologies:
                        p = self._docx_paragraph(doc)
                        self._docx_run(p, ", ".join(project.technologies), italic=True,
                                       size=s.body_size - 0.5, color=s.muted)
                    if project.description.strip():
                        self._docx_run(self._docx_paragraph(doc), project.description)
                    for item in project.responsibilities + project.achievements:
                        self._docx_bullet(doc, item)
                    if project.results.strip():
                        self._docx_bullet(doc, f"Result: {project.results}")
            elif key == "education":
                for edu in resume.education:
                    p = self._docx_paragraph(doc, space_before=3)
                    self._docx_run(p, " — ".join(x for x in (edu.degree, edu.institution) if x),
                                   bold=True)
                    meta = " | ".join(x for x in (
                        edu.location, date_range(edu.start_date, edu.end_date),
                        f"GPA/Score: {edu.gpa}" if edu.gpa else "") if x)
                    if meta:
                        self._docx_run(self._docx_paragraph(doc), meta, italic=True,
                                       size=s.body_size - 0.5, color=s.muted)
            elif key == "certifications":
                for cert in resume.certifications:
                    text = " — ".join(x for x in (cert.name, cert.issuer) if x)
                    if cert.date:
                        text += f" ({cert.date})"
                    if cert.credential:
                        text += f" | {cert.credential}"
                    self._docx_bullet(doc, text)
            elif key == "achievements":
                for ach in resume.achievements:
                    self._docx_bullet(doc, f"{ach.title}: {ach.description}"
                                      if ach.description else ach.title)

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()
